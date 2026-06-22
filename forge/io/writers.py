from __future__ import annotations

import csv
import json
import sqlite3
from typing import Any

from forge.core.plugin import Registry, Writer


@Registry.register_writer
class CSVWriter(Writer):
    format = "csv"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        if not records:
            with open(destination, "w") as f:
                pass
            return
        with open(destination, "w", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=records[0].keys())
            writer.writeheader()
            writer.writerows(records)


@Registry.register_writer
class JSONWriter(Writer):
    format = "json"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        with open(destination, "w") as f:
            json.dump(records, f, indent=kwargs.get("indent", 2), default=str)


@Registry.register_writer
class NDJSONWriter(Writer):
    format = "ndjson"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        with open(destination, "w") as f:
            for record in records:
                f.write(json.dumps(record, default=str) + "\n")


@Registry.register_writer
class XLSXWriter(Writer):
    format = "xlsx"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl required for XLSX. Install: pip install forge[xlsx]")
        wb = openpyxl.Workbook()
        ws = wb.active
        if records:
            ws.append(list(records[0].keys()))
            for r in records:
                ws.append(list(r.values()))
        wb.save(destination)


@Registry.register_writer
class YAMLWriter(Writer):
    format = "yaml"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        try:
            import yaml
        except ImportError:
            raise ImportError("pyyaml required for YAML. Install: pip install forge[yaml]")
        with open(destination, "w") as f:
            yaml.dump(records, f, default_flow_style=False, sort_keys=False)


@Registry.register_writer
class SQLiteWriter(Writer):
    format = "sqlite"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        table = kwargs.get("table", "data")
        conn = sqlite3.connect(destination)
        if records:
            cols = list(records[0].keys())
            placeholders = ",".join("?" for _ in cols)
            col_list = ",".join(f"[{c}]" for c in cols)
            conn.execute(f"CREATE TABLE IF NOT EXISTS [{table}] ({col_list})")
            for r in records:
                conn.execute(
                    f"INSERT INTO [{table}] ({col_list}) VALUES ({placeholders})",
                    [r.get(c) for c in cols],
                )
        conn.commit()
        conn.close()


@Registry.register_writer
class ParquetWriter(Writer):
    format = "parquet"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        try:
            import pandas as pd
            import pyarrow
        except ImportError:
            raise ImportError("pyarrow+pandas required for Parquet. Install: pip install forge[parquet]")
        df = pd.DataFrame(records)
        df.to_parquet(destination, index=False)


@Registry.register_writer
class XMLWriter(Writer):
    format = "xml"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        root_tag = kwargs.get("root_tag", "data")
        item_tag = kwargs.get("item_tag", "record")
        lines = [f'<?xml version="1.0" encoding="utf-8"?>', f"<{root_tag}>"]
        for r in records:
            lines.append(f"  <{item_tag}>")
            for k, v in r.items():
                val = str(v) if v is not None else ""
                lines.append(f"    <{k}>{val}</{k}>")
            lines.append(f"  </{item_tag}>")
        lines.append(f"</{root_tag}>")
        with open(destination, "w") as f:
            f.write("\n".join(lines))


@Registry.register_writer
class MarkdownWriter(Writer):
    format = "md"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        if not records:
            with open(destination, "w") as f:
                pass
            return
        headers = list(records[0].keys())
        sep = "|" + "|".join("---" for _ in headers) + "|"
        header = "|" + "|".join(headers) + "|"
        lines = [header, sep]
        for r in records:
            lines.append("|" + "|".join(str(r.get(h, "")) for h in headers) + "|")
        with open(destination, "w") as f:
            f.write("\n".join(lines) + "\n")


@Registry.register_writer
class SQLWriter(Writer):
    format = "sql"

    def write(self, records: list[dict], destination: str, **kwargs) -> None:
        table = kwargs.get("table", "data")
        if not records:
            with open(destination, "w") as f:
                pass
            return
        cols = list(records[0].keys())
        col_list = ", ".join(cols)
        stmts = [f"CREATE TABLE IF NOT EXISTS {table} ({col_list});"]
        for r in records:
            vals = []
            for c in cols:
                v = r.get(c)
                if v is None:
                    vals.append("NULL")
                elif isinstance(v, (int, float)):
                    vals.append(str(v))
                else:
                    vals.append("'" + str(v).replace("'", "''") + "'")
            stmts.append(f"INSERT INTO {table} ({col_list}) VALUES ({', '.join(vals)});")
        with open(destination, "w") as f:
            f.write("\n".join(stmts) + "\n")
